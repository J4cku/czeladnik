"""Extracts both final workbooks and their drawings into the question catalog."""
import hashlib, json, re, unicodedata
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "lib" / "questions.json"
IMAGE_OUT = ROOT / "public" / "question-images"

# sheet title -> (id, kind, short label, description)
SHEETS = {
    "1. Rachunkowość":      ("rachunkowosc", "abc", "Rachunkowość", "Obliczenia optyczne, moce soczewek, pryzmaty, ekwiwalent sferyczny."),
    "2. Dokumentacja":      ("dokumentacja", "abc", "Dokumentacja", "Dokumenty, podatki, ewidencja, formy prawne działalności."),
    "3. Rysunek zawodowy":  ("rysunek", "open", "Rysunek zawodowy", "Anatomia oka, symbole i przekroje soczewek, bieg promieni."),
    "4. BHP i ppoż.":       ("bhp", "abc", "BHP i ppoż.", "Bezpieczeństwo pracy, pierwsza pomoc, ochrona przeciwpożarowa."),
    "5. Ochrona środowiska":("srodowisko", "abc", "Ochrona środowiska", "Odpady, emisje, ekologia i gospodarka obiegu zamkniętego."),
    "6. Prawo pracy":       ("prawo-pracy", "abc", "Prawo pracy", "Kodeks pracy, umowy, pracownicy młodociani, urlopy."),
    "7. Działalność gosp.": ("dzialalnosc", "abc", "Działalność gospodarcza", "Ekonomia, podatki, ZUS, prowadzenie firmy."),
    "USTNY 1. Technologia": ("ustny-technologia", "open", "Technologia", "Pytania ustne: oko, refrakcja, wykonanie i centracja okularów."),
    "USTNY 2. Materiałozn.":("ustny-materialy", "open", "Materiałoznawstwo", "Pytania ustne: materiały soczewek, opraw, powłoki."),
    "USTNY 3. Maszynozn.":  ("ustny-maszyny", "open", "Maszynoznawstwo", "Pytania ustne: narzędzia i urządzenia pracowni optycznej."),
    "Zadania praktyczne":   ("praktyka", "task", "Zadania praktyczne", "Zadania na część praktyczną egzaminu wraz z czasem wykonania."),
}

MISTRZ_SHEETS = {
    title: (cat_id, "abc" if cat_id == "rysunek" else kind, label, description)
    for title, (cat_id, kind, label, description) in SHEETS.items()
    if kind != "task"
}
MISTRZ_SHEETS.update({
    "8. Psychologia i pedag.": ("psychologia", "abc", "Psychologia i pedagogika", "Psychologia, pedagogika i praca z uczniami."),
    "9. Metodyka nauczania": ("metodyka", "abc", "Metodyka nauczania", "Organizacja i metody nauczania zawodu."),
})

LEVELS = {
    "czeladnik": ("Czeladnik_optyk_pytania_odpowiedzi.xlsx", SHEETS),
    "mistrz": ("Mistrz_optyk_pytania_odpowiedzi.xlsx", MISTRZ_SHEETS),
}

DIFFICULTY = {
    "Ł": "latwe",
    "ŁATWY": "latwe",
    "LATWY": "latwe",
    "Ś": "srednie",
    "S": "srednie",
    "ŚREDNI": "srednie",
    "SREDNI": "srednie",
    "T": "trudne",
    "TRUDNY": "trudne",
}


def clean(v):
    if v is None:
        return None
    s = str(v).replace(" ", " ").strip()
    s = unicodedata.normalize("NFC", s)
    s = re.sub(r"\s+", " ", s)
    return s or None


def as_index(v):
    """'1.0' -> 1"""
    s = clean(v)
    if s is None:
        return None
    try:
        return int(float(s))
    except ValueError:
        return None


def normalize_header(value):
    value = clean(value)
    if value is None:
        return None
    value = unicodedata.normalize("NFKD", value.casefold().replace("ł", "l"))
    value = "".join(character for character in value if not unicodedata.combining(character))
    return {
        "numer pytania": "nr pytania",
        "id": "nr pytania",
        "poziom": "trudnosc",
        "poziom trudnosci": "trudnosc",
    }.get(value, value)


def extract_workbook(source, level, sheets, seen_ids, assets):
    wb = openpyxl.load_workbook(source, data_only=True)
    categories, questions = [], []

    for ws in wb.worksheets:
        if ws.title not in sheets:
            raise SystemExit(f"Nieznany arkusz: {ws.title}")
        cat_id, kind, label, description = sheets[ws.title]
        if level == "mistrz":
            cat_id = f"mistrz-{cat_id}"
        header = [normalize_header(c.value) for c in ws[1]]
        images = {}
        for image in ws._images:
            row_number = image.anchor._from.row + 1
            if row_number in images:
                raise SystemExit(f"Wiele rysunków: {ws.title}, wiersz {row_number}")
            content = image._data()
            if not content.startswith(b"\x89PNG\r\n\x1a\n"):
                raise SystemExit(f"Rysunek nie jest PNG: {ws.title}, wiersz {row_number}")
            filename = hashlib.sha256(content).hexdigest() + ".png"
            assets[filename] = content
            images[row_number] = {
                "src": f"/question-images/{filename}",
                "width": image.width,
                "height": image.height,
            }
        count = 0

        for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            cells = dict(zip(header, row))
            prompt = clean(cells.get("pytanie") or cells.get("zadanie"))
            if not prompt:
                continue
            source_nr = as_index(cells.get("nr"))
            official = as_index(cells.get("nr pytania"))
            nr = source_nr or count + 1
            q = {"category": cat_id, "kind": kind, "nr": nr, "prompt": prompt}

            if kind == "abc":
                options = [clean(cells.get(k)) for k in ("a", "b", "c")]
                if any(o is None for o in options):
                    raise SystemExit(f"Brak odpowiedzi: {cat_id}, wiersz {row_number}")
                letter = (clean(cells.get("odpowiedz")) or "").upper()
                if letter not in ("A", "B", "C"):
                    raise SystemExit(f"Zła litera odpowiedzi: {cat_id}, wiersz {row_number}: {letter!r}")
                q["options"] = options
                q["answer"] = "ABC".index(letter)
            elif kind == "open":
                q["answer"] = (
                    clean(cells.get("odpowiedz z klucza"))
                    or clean(cells.get("odpowiedz"))
                    or "Brak opisu odpowiedzi."
                )
            else:
                q["time"] = clean(cells.get("czas"))

            if official:
                qid = f"{cat_id}-{official}"
            else:
                identity = json.dumps(
                    {key: value for key, value in q.items() if key != "nr"},
                    ensure_ascii=False,
                    sort_keys=True,
                    separators=(",", ":"),
                )
                digest = hashlib.sha256(identity.encode()).hexdigest()[:16]
                qid = f"{cat_id}-h-{digest}"
            if qid in seen_ids:
                raise SystemExit(f"Duplikat id: {qid} (wiersz {row_number})")
            seen_ids.add(qid)
            q = {"id": qid, **q, "level": level}

            difficulty = clean(cells.get("trudnosc")) or ""
            diff = DIFFICULTY.get(difficulty.upper())
            if diff:
                q["difficulty"] = diff
            if official:
                q["officialNr"] = official
            if row_number in images:
                q["image"] = images.pop(row_number)

            questions.append(q)
            count += 1

        if images:
            raise SystemExit(f"Rysunek bez pytania: {ws.title}, wiersze {sorted(images)}")
        categories.append({"id": cat_id, "kind": kind, "label": label,
                           "description": description, "count": count, "level": level})

    wb.close()
    return categories, questions


def main():
    categories, questions = [], []
    seen_ids, assets = set(), {}
    for level, (filename, sheets) in LEVELS.items():
        level_categories, level_questions = extract_workbook(
            ROOT / "data" / filename, level, sheets, seen_ids, assets,
        )
        categories.extend(level_categories)
        questions.extend(level_questions)

    OUT.parent.mkdir(exist_ok=True)
    OUT.write_text(json.dumps({"categories": categories, "questions": questions},
                              ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    IMAGE_OUT.mkdir(parents=True, exist_ok=True)
    for filename, content in assets.items():
        (IMAGE_OUT / filename).write_bytes(content)

    for c in categories:
        print(f"{c['kind']:5} {c['id']:20} {c['count']:4}")
    print("razem:", len(questions))
    print("rysunki:", sum("image" in question for question in questions), "pliki:", len(assets))


if __name__ == "__main__":
    main()
