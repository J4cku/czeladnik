import hashlib
import json
import re
import unittest
import unicodedata
from collections import Counter
from pathlib import Path

import openpyxl
from PIL import Image


ROOT = Path(__file__).resolve().parent.parent
SOURCES = {
    "czeladnik": ROOT / "data/Czeladnik_optyk_pytania_odpowiedzi.xlsx",
    "mistrz": ROOT / "data/Mistrz_optyk_pytania_odpowiedzi.xlsx",
}


def catalog(level=None):
    data = json.loads((ROOT / "lib" / "questions.json").read_text(encoding="utf-8"))
    if level:
        data = {
            key: [entry for entry in entries if entry.get("level", "czeladnik") == level]
            for key, entries in data.items()
        }
    return data


def clean(value):
    return re.sub(
        r"\s+",
        " ",
        unicodedata.normalize("NFC", str(value).replace("\xa0", " ").strip()),
    )


class QuestionsDataTest(unittest.TestCase):
    def test_question_content_matches_checked_workbook_snapshot(self):
        data = catalog("czeladnik")
        for category in data["categories"]:
            category.pop("level", None)
        for question in data["questions"]:
            question.pop("difficulty", None)
            question.pop("officialNr", None)
            question.pop("level", None)
            question.pop("image", None)

        content = json.dumps(
            data,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        ).encode()

        self.assertEqual(
            "b0db7f3b773f78c5f78b6076a5fc463becde127743f3a490b9a185b243a5618b",
            hashlib.sha256(content).hexdigest(),
        )

    def test_updated_categories_include_all_source_difficulties(self):
        data = json.loads((ROOT / "lib" / "questions.json").read_text(encoding="utf-8"))

        expected = {
            "rachunkowosc": Counter(latwe=20, srednie=20, trudne=20),
            "dokumentacja": Counter(latwe=39, srednie=26, trudne=19),
            "rysunek": Counter(latwe=20, srednie=19, trudne=20),
            "bhp": Counter(latwe=30, srednie=30, trudne=19),
            "srodowisko": Counter(latwe=30, srednie=27, trudne=21),
            "prawo-pracy": Counter(latwe=30, srednie=30, trudne=20),
            "dzialalnosc": Counter(latwe=30, srednie=30, trudne=20),
            "ustny-technologia": Counter(latwe=27, srednie=27, trudne=24),
            "ustny-materialy": Counter(latwe=22, srednie=22, trudne=21),
            "ustny-maszyny": Counter(latwe=13, srednie=13, trudne=13),
            "mistrz-rachunkowosc": Counter(latwe=21, srednie=20, trudne=19),
            "mistrz-dokumentacja": Counter(latwe=30, srednie=30, trudne=20),
            "mistrz-rysunek": Counter(latwe=20, srednie=20, trudne=20),
            "mistrz-bhp": Counter(latwe=30, srednie=30, trudne=20),
            "mistrz-srodowisko": Counter(latwe=30, srednie=30, trudne=20),
            "mistrz-prawo-pracy": Counter(latwe=30, srednie=30, trudne=18),
            "mistrz-dzialalnosc": Counter(latwe=30, srednie=27, trudne=19),
            "mistrz-psychologia": Counter(latwe=30, srednie=30, trudne=20),
            "mistrz-metodyka": Counter(latwe=30, srednie=30, trudne=20),
            "mistrz-ustny-technologia": Counter(latwe=26, srednie=26, trudne=26),
            "mistrz-ustny-materialy": Counter(latwe=20, srednie=20, trudne=20),
            "mistrz-ustny-maszyny": Counter(latwe=13, srednie=13, trudne=13),
        }

        for category, difficulties in expected.items():
            with self.subTest(category=category):
                actual = Counter(
                    question["difficulty"]
                    for question in data["questions"]
                    if question["category"] == category and "difficulty" in question
                )
                self.assertEqual(difficulties, actual)

    def test_updated_categories_include_source_question_numbers(self):
        data = json.loads((ROOT / "lib" / "questions.json").read_text(encoding="utf-8"))

        expected = {
            "rachunkowosc": 60,
            "dokumentacja": 73,
            "rysunek": 59,
            "bhp": 79,
            "srodowisko": 78,
            "prawo-pracy": 80,
            "dzialalnosc": 80,
            "ustny-technologia": 78,
            "ustny-materialy": 63,
            "ustny-maszyny": 39,
            "mistrz-rachunkowosc": 60,
            "mistrz-dokumentacja": 80,
            "mistrz-rysunek": 60,
            "mistrz-bhp": 80,
            "mistrz-srodowisko": 80,
            "mistrz-prawo-pracy": 78,
            "mistrz-dzialalnosc": 76,
            "mistrz-psychologia": 80,
            "mistrz-metodyka": 80,
            "mistrz-ustny-technologia": 76,
            "mistrz-ustny-materialy": 60,
            "mistrz-ustny-maszyny": 39,
        }

        for category, count in expected.items():
            with self.subTest(category=category):
                actual = sum(
                    "officialNr" in question
                    for question in data["questions"]
                    if question["category"] == category
                )
                self.assertEqual(count, actual)

    def test_category_counts_match_checked_workbook(self):
        data = json.loads((ROOT / "lib" / "questions.json").read_text(encoding="utf-8"))

        expected = {
            "rachunkowosc": 60,
            "dokumentacja": 84,
            "rysunek": 59,
            "bhp": 79,
            "srodowisko": 78,
            "prawo-pracy": 80,
            "dzialalnosc": 80,
            "ustny-technologia": 78,
            "ustny-materialy": 65,
            "ustny-maszyny": 39,
            "praktyka": 30,
            "mistrz-rachunkowosc": 60,
            "mistrz-dokumentacja": 80,
            "mistrz-rysunek": 60,
            "mistrz-bhp": 80,
            "mistrz-srodowisko": 80,
            "mistrz-prawo-pracy": 78,
            "mistrz-dzialalnosc": 78,
            "mistrz-psychologia": 80,
            "mistrz-metodyka": 80,
            "mistrz-ustny-technologia": 78,
            "mistrz-ustny-materialy": 60,
            "mistrz-ustny-maszyny": 39,
        }

        actual = {category["id"]: category["count"] for category in data["categories"]}
        self.assertEqual(expected, actual)
        self.assertEqual(1585, len(data["questions"]))

    def test_catalog_levels_match_source_totals(self):
        data = catalog()
        self.assertEqual(
            Counter(czeladnik=732, mistrz=853),
            Counter(question.get("level") for question in data["questions"]),
        )
        self.assertEqual(
            Counter(czeladnik=11, mistrz=12),
            Counter(category.get("level") for category in data["categories"]),
        )
        levels = {category["id"]: category["level"] for category in data["categories"]}
        for question in data["questions"]:
            self.assertEqual(levels[question["category"]], question["level"])

    def test_apprentice_question_ids_remain_stable(self):
        ids = sorted(question["id"] for question in catalog("czeladnik")["questions"])
        self.assertEqual(732, len(ids))
        self.assertEqual(
            "724a060527ca5da1772080c07b21cd2a9317668a7c5a69509d45c0418b29a085",
            hashlib.sha256("\n".join(ids).encode()).hexdigest(),
        )

    def test_master_questions_use_namespaced_ids_and_drawing_answers(self):
        questions = catalog("mistrz")["questions"]
        self.assertEqual(853, len(questions))
        for question in questions:
            self.assertTrue(question["id"].startswith("mistrz-"))
            self.assertTrue(question["category"].startswith("mistrz-"))
            self.assertNotEqual("task", question["kind"])
            if question["category"] == "mistrz-rysunek":
                self.assertEqual("abc", question["kind"])
                self.assertEqual(3, len(question["options"]))
                self.assertIn(question["answer"], (0, 1, 2))

    def test_master_business_keeps_two_ungraded_questions(self):
        questions = [
            question for question in catalog("mistrz")["questions"]
            if question["category"] == "mistrz-dzialalnosc"
            and "difficulty" not in question
        ]
        self.assertEqual(2, len(questions))
        self.assertTrue(all(question["prompt"] and question["options"] for question in questions))

    def test_checked_workbooks_match_final_source_files(self):
        expected = {
            "czeladnik": "e2f8a910022a3164b682951af541e1ffa96163f3bf9a46b880ca9aa1e28382ee",
            "mistrz": "4f73b2fc06dfd6047adaae43f848e7113e1920dbda257a6c9043f567f82e5265",
        }
        for level, source in SOURCES.items():
            with self.subTest(level=level):
                self.assertTrue(source.exists())
                self.assertEqual(expected[level], hashlib.sha256(source.read_bytes()).hexdigest())

    def test_master_oral_answers_use_normalized_answer_key_headers(self):
        questions = catalog("mistrz")["questions"]
        self.assertEqual(853, len(questions))
        generated = {(q["category"], q["prompt"]): q["answer"] for q in questions}
        workbook = openpyxl.load_workbook(SOURCES["mistrz"], data_only=True, read_only=True)
        sheets = {
            "USTNY 1. Technologia": "mistrz-ustny-technologia",
            "USTNY 2. Materiałozn.": "mistrz-ustny-materialy",
            "USTNY 3. Maszynozn.": "mistrz-ustny-maszyny",
        }
        for sheet_name, category in sheets.items():
            sheet = workbook[sheet_name]
            header = [clean(cell.value).lower() for cell in sheet[1]]
            for row in sheet.iter_rows(min_row=2, values_only=True):
                cells = dict(zip(header, row))
                if not cells["pytanie"]:
                    continue
                with self.subTest(sheet=sheet_name, prompt=cells["pytanie"]):
                    self.assertEqual(
                        clean(cells["odpowiedz z klucza"] or cells["odpowiedź"]),
                        generated[(category, clean(cells["pytanie"]))],
                    )

    def test_drawing_images_match_their_source_rows_and_assets(self):
        data = catalog()
        self.assertEqual(
            Counter(czeladnik=59, mistrz=60),
            Counter(question.get("level") for question in data["questions"] if "image" in question),
        )
        generated = {q["id"]: q for q in data["questions"]}
        for level, source in SOURCES.items():
            workbook = openpyxl.load_workbook(source, data_only=True)
            for sheet in workbook.worksheets:
                for image in sheet._images:
                    row_number = image.anchor._from.row + 1
                    prompt_column = next(cell.column for cell in sheet[1] if cell.value == "Pytanie")
                    prompt = clean(sheet.cell(row_number, prompt_column).value)
                    official_column = next(
                        cell.column for cell in sheet[1] if cell.value in ("Nr pytania", "ID")
                    )
                    official = int(sheet.cell(row_number, official_column).value)
                    category = "rysunek" if level == "czeladnik" else "mistrz-rysunek"
                    with self.subTest(level=level, row=row_number):
                        question = generated[f"{category}-{official}"]
                        self.assertEqual(prompt, question["prompt"])
                        asset = question["image"]
                        self.assertEqual({"src", "width", "height"}, set(asset))
                        self.assertEqual((image.width, image.height), (asset["width"], asset["height"]))
                        image_data = image._data()
                        digest = hashlib.sha256(image_data).hexdigest()
                        self.assertEqual(f"/question-images/{digest}.png", asset["src"])
                        path = ROOT / "public" / asset["src"].lstrip("/")
                        self.assertEqual(image_data, path.read_bytes())
                        with Image.open(path) as exported:
                            self.assertEqual("PNG", exported.format)
                            self.assertEqual((asset["width"], asset["height"]), exported.size)

    def test_question_ids_are_unique(self):
        data = json.loads((ROOT / "lib" / "questions.json").read_text(encoding="utf-8"))
        ids = [question["id"] for question in data["questions"]]

        self.assertEqual(len(ids), len(set(ids)))

    def test_question_ids_use_official_numbers_or_content_hashes(self):
        data = json.loads((ROOT / "lib" / "questions.json").read_text(encoding="utf-8"))

        for question in data["questions"]:
            with self.subTest(question=question["id"]):
                if "officialNr" in question:
                    self.assertEqual(
                        f'{question["category"]}-{question["officialNr"]}',
                        question["id"],
                    )
                else:
                    identity = {
                        key: value
                        for key, value in question.items()
                        if key not in {"id", "nr", "difficulty", "officialNr", "level", "image"}
                    }
                    digest = hashlib.sha256(
                        json.dumps(
                            identity,
                            ensure_ascii=False,
                            sort_keys=True,
                            separators=(",", ":"),
                        ).encode()
                    ).hexdigest()[:16]
                    self.assertEqual(
                        f'{question["category"]}-h-{digest}',
                        question["id"],
                    )

    def test_oral_answers_match_the_checked_answer_key_columns(self):
        data = json.loads((ROOT / "lib" / "questions.json").read_text(encoding="utf-8"))
        generated = {
            (question["category"], question["prompt"]): question["answer"]
            for question in data["questions"]
            if question["category"].startswith("ustny-")
        }
        workbook = openpyxl.load_workbook(
            ROOT / "data" / "Czeladnik_optyk_pytania_odpowiedzi.xlsx",
            data_only=True,
            read_only=True,
        )
        sheets = {
            "USTNY 1. Technologia": "ustny-technologia",
            "USTNY 2. Materiałozn.": "ustny-materialy",
            "USTNY 3. Maszynozn.": "ustny-maszyny",
        }

        for sheet_name, category in sheets.items():
            sheet = workbook[sheet_name]
            header = [clean(cell.value) if cell.value is not None else None for cell in sheet[1]]
            for row in sheet.iter_rows(min_row=2, values_only=True):
                cells = dict(zip(header, row))
                prompt = cells.get("Pytanie")
                answer = cells.get("Odpowiedź z klucza") or cells.get("Odpowiedz z klucza")
                if prompt and answer:
                    with self.subTest(sheet=sheet_name, prompt=prompt):
                        self.assertEqual(clean(answer), generated[(category, clean(prompt))])


if __name__ == "__main__":
    unittest.main()
